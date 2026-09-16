import os
from numbers import Real
from datetime import datetime
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


class ExcelExporter:
    """Genera un reporte de Excel con KPIs y datos financieros."""

    CURRENCY_FORMAT = "$#,##0.0000"
    PERCENT_FORMAT = "0.00%"
    DATE_FORMAT = "yyyy-mm-dd"
    CURRENCY_COLUMNS = {"tipo_cambio", "ma_30d", "ma_90d"}
    PERCENT_COLUMNS = {
        "variacion_diaria_pct",
        "variacion_mensual_pct",
        "volatilidad_30d",
    }
    REQUIRED_KPIS = {
        "fecha_corte",
        "tc_actual",
        "var_diaria_pct",
        "tc_max_anio",
        "tc_min_anio",
        "tc_prom_anio",
    }

    @staticmethod
    def export_to_excel(
        df_usd: pd.DataFrame,
        kpis: dict,
        output_path: str = "data/reporte_financiero.xlsx",
    ) -> str:
        """Guarda el DataFrame y los KPIs en un archivo .xlsx y devuelve su ruta."""
        if not isinstance(df_usd, pd.DataFrame):
            raise TypeError("df_usd debe ser un pandas.DataFrame")
        if df_usd.empty:
            raise ValueError("df_usd no puede estar vacio")
        if df_usd.columns.duplicated().any():
            raise ValueError("df_usd no puede contener columnas duplicadas")
        required_columns = {"fecha", "tipo_cambio"}
        missing_columns = required_columns - set(df_usd.columns)
        if missing_columns:
            raise ValueError(
                f"Faltan columnas requeridas: {', '.join(sorted(missing_columns))}"
            )
        if not pd.api.types.is_datetime64_any_dtype(df_usd["fecha"]):
            raise TypeError("La columna fecha debe tener tipo datetime")
        if df_usd["fecha"].isna().any():
            raise ValueError("La columna fecha no puede contener valores nulos")
        if not pd.api.types.is_numeric_dtype(df_usd["tipo_cambio"]):
            raise TypeError("La columna tipo_cambio debe ser numerica")
        missing_kpis = ExcelExporter.REQUIRED_KPIS - set(kpis)
        if missing_kpis:
            missing = ", ".join(sorted(missing_kpis))
            raise ValueError(f"Faltan KPIs requeridos: {missing}")
        for name in ExcelExporter.REQUIRED_KPIS - {"fecha_corte"}:
            if not isinstance(kpis[name], Real):
                raise TypeError(f"El KPI {name} debe ser numerico")

        output = Path(output_path)
        output.parent.mkdir(parents=True, exist_ok=True)
        workbook = openpyxl.Workbook()
        summary = workbook.active
        if summary is None:
            raise RuntimeError("No se pudo crear la hoja de resumen")
        summary.title = "Resumen Ejecutivo"

        title_font = Font(name="Calibri", size=16, bold=True, color="1F4E79")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        center = Alignment(horizontal="center", vertical="center")
        border = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9"),
        )

        summary["A1"] = "REPORTE FINANCIERO AUTOMATIZADO - BANXICO"
        summary["A1"].font = title_font
        summary["A2"] = f"Fecha de corte: {kpis['fecha_corte']}"
        summary["A2"].font = Font(italic=True, color="595959")
        summary.merge_cells("A1:C1")

        headers = ["Indicador Financiero", "Valor", "Unidad / Formato"]
        for column, header in enumerate(headers, 1):
            cell = summary.cell(row=4, column=column, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center
            cell.border = border

        rows = [
            ("Tipo de Cambio Actual (USD/MXN)", kpis["tc_actual"], ExcelExporter.CURRENCY_FORMAT),
            ("Variación Diaria %", kpis["var_diaria_pct"] / 100, ExcelExporter.PERCENT_FORMAT),
            ("Máximo del Año", kpis["tc_max_anio"], ExcelExporter.CURRENCY_FORMAT),
            ("Mínimo del Año", kpis["tc_min_anio"], ExcelExporter.CURRENCY_FORMAT),
            ("Promedio del Año", kpis["tc_prom_anio"], ExcelExporter.CURRENCY_FORMAT),
        ]
        for row_number, (label, value, number_format) in enumerate(rows, 5):
            summary.cell(row=row_number, column=1, value=label)
            value_cell = summary.cell(row=row_number, column=2, value=float(value))
            value_cell.number_format = number_format
            summary.cell(row=row_number, column=3, value=number_format)
            for column in range(1, 4):
                summary.cell(row=row_number, column=column).border = border

        summary.column_dimensions["A"].width = 38
        summary.column_dimensions["B"].width = 18
        summary.column_dimensions["C"].width = 20

        data_sheet = workbook.create_sheet("Histórico")
        data = df_usd.copy()
        for column, name in enumerate(data.columns, 1):
            cell = data_sheet.cell(row=1, column=column, value=name)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center
            cell.border = border
        for row_number, values in enumerate(data.itertuples(index=False, name=None), 2):
            for column, value in enumerate(values, 1):
                column_name = data.columns[column - 1]
                cell_value = None if pd.isna(value) else value
                if cell_value is not None and column_name in ExcelExporter.PERCENT_COLUMNS:
                    cell_value = float(cell_value) / 100
                cell = data_sheet.cell(row=row_number, column=column, value=cell_value)
                cell.border = border
                if column_name == "fecha" and cell_value is not None:
                    cell.number_format = ExcelExporter.DATE_FORMAT
                elif column_name in ExcelExporter.CURRENCY_COLUMNS:
                    cell.number_format = ExcelExporter.CURRENCY_FORMAT
                elif column_name in ExcelExporter.PERCENT_COLUMNS:
                    cell.number_format = ExcelExporter.PERCENT_FORMAT
                elif isinstance(cell_value, Real) and not isinstance(cell_value, bool):
                    cell.number_format = "0.0000"

        if data.columns.size:
            data_sheet.freeze_panes = "A2"
            data_sheet.auto_filter.ref = data_sheet.dimensions
            for column_number, column_name in enumerate(data.columns, 1):
                values = [str(column_name)] + [
                    str(value) for value in data.iloc[:, column_number - 1].dropna()
                ]
                width = min(max(len(value) for value in values) + 2, 24)
                data_sheet.column_dimensions[get_column_letter(column_number)].width = width

        chart_columns = [
            index
            for index, column_name in enumerate(data.columns, 1)
            if column_name in {"tipo_cambio", "ma_30d", "ma_90d"}
        ]
        if len(data) > 1 and chart_columns:
            chart = LineChart()
            chart.title = "Evolución del tipo de cambio"
            chart.y_axis.title = "USD/MXN"
            chart.x_axis.title = "Fecha"
            for column_number in chart_columns:
                chart.add_data(
                    Reference(
                        data_sheet,
                        min_col=column_number,
                        max_col=column_number,
                        min_row=1,
                        max_row=len(data) + 1,
                    ),
                    titles_from_data=True,
                )
            chart.set_categories(Reference(data_sheet, min_col=1, min_row=2, max_row=len(data) + 1))
            chart.height = 8
            chart.width = 18
            data_sheet.add_chart(chart, "E2")

        try:
            workbook.save(output)
        except PermissionError:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            output = output.with_name(f"{output.stem}_{timestamp}{output.suffix}")
            workbook.save(output)
            print(f"El archivo original estaba bloqueado; se genero una copia en: {output}")
        return os.fspath(output)
